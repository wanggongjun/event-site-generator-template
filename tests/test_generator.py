"""Generator regression tests; edit XLSX XML fixtures without another authoring tool."""
from pathlib import Path
from datetime import date
import io
import json
import shutil
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from generate import ConfigError, generate, load_config, event_facts, resource_destination, public_config
NS={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def mutate_workbook(source, destination, sheet_name, updates=None, duplicate_row=None):
    """Patch only requested XML fixture cells; no spreadsheet app side effects."""
    with ZipFile(source) as z:
        files={name:z.read(name) for name in z.namelist()}
    workbook=ET.fromstring(files['xl/workbook.xml'])
    rid=None
    for sheet in workbook.findall('s:sheets/s:sheet',NS):
        if sheet.attrib['name']==sheet_name:
            rid=sheet.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']
    relationships=ET.fromstring(files['xl/_rels/workbook.xml.rels'])
    target=next(r.attrib['Target'] for r in relationships if r.attrib['Id']==rid)
    filename=target.lstrip('/') if target.startswith('/') else 'xl/'+target
    tree=ET.fromstring(files[filename]);data=tree.find('s:sheetData',NS)
    for cell in data.findall('s:row/s:c',NS):
        if cell.attrib['r'] in (updates or {}):
            value=updates[cell.attrib['r']]
            for child in list(cell):
                cell.remove(child)
            cell.attrib['t']='inlineStr'
            ET.SubElement(ET.SubElement(cell,'{'+NS['s']+'}is'),'{'+NS['s']+'}t').text=str(value)
    if duplicate_row:
        old=next(r for r in data if r.attrib['r']==str(duplicate_row))
        row=ET.fromstring(ET.tostring(old));number=max(int(r.attrib['r']) for r in data)+1
        row.attrib['r']=str(number)
        for c in row:
            if 'r' in c.attrib:
                c.attrib['r']=''.join(ch for ch in c.attrib['r'] if ch.isalpha())+str(number)
        data.append(row)
    files[filename]=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
    with ZipFile(destination,'w',ZIP_DEFLATED) as z:
        for name,content in files.items():
            z.writestr(name,content)


class GeneratorTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(prefix='event-generator-test-')
        self.base=Path(self.tmp.name)
        self.input=self.base/'input';shutil.copytree(ROOT/'input',self.input)
        self.workbook=self.input/'fictional-conference.xlsx'
        self.output=self.base/'event'
        self.template=self.base/'template';self.template.mkdir()
        (self.template/'package.json').write_text('{"scripts":{"start":"node backend/index.js"}}')
        (self.template/'backend').mkdir();(self.template/'backend/index.js').write_text('// runtime placeholder')
        (self.template/'backend/.env.example').write_text('PORT=3000\n')
        (self.template/'backend/data').mkdir();(self.template/'backend/data/SHOULD_NOT_COPY').write_text('business')
        (self.template/'node_modules').mkdir();(self.template/'node_modules/SHOULD_NOT_COPY').write_text('dependency')

    def tearDown(self):
        self.tmp.cleanup()

    def change(self,sheet,updates=None,duplicate_row=None):
        new=self.input/'changed.xlsx';mutate_workbook(self.workbook,new,sheet,updates,duplicate_row)
        return new

    def run_generate(self,workbook=None):
        return generate(workbook or self.workbook,self.output,self.template)

    def test_generation_independent_config_assets_and_materials(self):
        _,manifest=self.run_generate()
        self.assertEqual(manifest['eventSlug'],'aurora-research-forum-2027')
        config=json.loads((self.output/'config.json').read_text())
        self.assertEqual(config['files'],{'maxFileBytes':20971520,'maxAttachments':3,'allowedExtensions':['pdf','docx','pptx']})
        self.assertEqual(config['recordings'],[])
        self.assertTrue((self.output/'frontend/public/assets/hero.svg').exists())
        for mirror in ('frontend/public/config.json','frontend/src/config.generated.json'):
            self.assertEqual(json.loads((self.output/mirror).read_text()),public_config(config))
            self.assertNotIn('readiness', json.loads((self.output/mirror).read_text()))
            self.assertNotIn('sourceNotes', json.loads((self.output/mirror).read_text()))
        for name in ('poster.svg','conference-guide.docx','conference-guide.html','social-posts.txt'):
            self.assertTrue((self.output/'materials'/name).exists())
        self.assertFalse((self.output/'backend/data').exists())
        self.assertFalse((self.output/'node_modules').exists())

    def test_arbitrary_event_names_and_escaped_svg_html(self):
        workbook=self.change('event',{'B3':'新名 & <研究> 未来会议','B4':'未来','B9':'星河大厅'})
        self.run_generate(workbook)
        cfg=json.loads((self.output/'config.json').read_text())
        self.assertEqual(cfg['event']['title'],'新名 & <研究> 未来会议')
        svg=(self.output/'frontend/public/assets/hero.svg').read_text()
        self.assertIn('新名 &amp; &lt;研究&gt; 未来会议',svg)
        self.assertIn('星河大厅',svg)
        self.assertNotIn('磐石',svg)

    def test_recordings_toggle_default_off_and_explicit_on(self):
        cfg,_=load_config(self.workbook);self.assertFalse(cfg['event']['recordingsEnabled']);self.assertEqual(cfg['recordings'],[])
        cfg,_=load_config(self.change('event',{'B13':'true'}))
        self.assertTrue(cfg['event']['recordingsEnabled']);self.assertEqual(len(cfg['recordings']),1)

    def test_invalid_config_never_touches_output(self):
        for updates in ({'B10':'0'},{'B13':'maybe'},{'B8':'Mars/Unknown'},{'B17':'2028-01-01T00:00:00'},{'B12':'javascript:alert(1)'},{'B20':'../../secret.png'},{'A20':'smtp.password'}):
            with self.subTest(updates=updates):
                bad=self.change('event',updates)
                with self.assertRaises(ConfigError):
                    self.run_generate(bad)
                self.assertFalse(self.output.exists())

    def test_missing_required_and_duplicate_rows_rejected(self):
        for sheet,kwargs in [('event',{'updates':{'B3':''}}),('event',{'duplicate_row':3}),('agenda',{'duplicate_row':2}),('faqs',{'duplicate_row':2})]:
            with self.subTest(sheet=sheet,kwargs=kwargs):
                with self.assertRaises(ConfigError):
                    self.run_generate(self.change(sheet,**kwargs))
                self.assertFalse(self.output.exists())

    def test_unknown_directory_and_slug_change_rejected(self):
        self.output.mkdir();(self.output/'important.txt').write_text('keep')
        with self.assertRaises(ConfigError):self.run_generate()
        self.assertEqual((self.output/'important.txt').read_text(),'keep')
        shutil.rmtree(self.output)
        self.run_generate()
        with self.assertRaises(ConfigError):self.run_generate(self.change('event',{'B2':'another-event'}))
        self.assertEqual(json.loads((self.output/'manifest.json').read_text())['eventSlug'],'aurora-research-forum-2027')

    def test_regeneration_preserves_database_uploads_credentials_and_unknown_files(self):
        self.run_generate()
        protected={'backend/data/event.sqlite':b'BUSINESS-STATE','backend/uploads/a.pdf':b'PRIVATE-UPLOAD','backend/.env':b'DATABASE_URL=private','uploads/b.pdf':b'STORED','operator-notes.txt':b'KEEP'}
        for rel,content in protected.items():
            p=self.output/rel;p.parent.mkdir(parents=True,exist_ok=True);p.write_bytes(content)
        self.run_generate(self.change('event',{'B3':'更新标题'}))
        for rel,content in protected.items():
            self.assertEqual((self.output/rel).read_bytes(),content)
        self.assertEqual(json.loads((self.output/'config.json').read_text())['event']['title'],'更新标题')

    def test_bad_regeneration_preserves_all_existing_files(self):
        self.run_generate()
        before={p.relative_to(self.output):p.read_bytes() for p in self.output.rglob('*') if p.is_file()}
        with self.assertRaises(ConfigError):self.run_generate(self.change('event',{'B6':'2029-01-01T09:00:00'}))
        self.assertEqual(before,{p.relative_to(self.output):p.read_bytes() for p in self.output.rglob('*') if p.is_file()})

    def test_repeat_generation_is_byte_identical(self):
        _,a=self.run_generate();_,b=self.run_generate()
        self.assertEqual(a,b)
        self.assertEqual(a['sha256'],b['sha256'])

    def test_shared_facts_across_all_materials(self):
        self.run_generate();cfg=json.loads((self.output/'config.json').read_text());facts=event_facts(cfg)
        svg=(self.output/'materials/poster.svg').read_text();hero=(self.output/'frontend/public/assets/hero.svg').read_text()
        posts=(self.output/'materials/social-posts.txt').read_text();guide=(self.output/'materials/conference-guide.html').read_text()
        with ZipFile(self.output/'materials/conference-guide.docx') as z:doc=z.read('word/document.xml').decode()
        for text in (facts['title'],facts['dates'],facts['location']):
            for artifact in (svg,posts,guide,doc):self.assertIn(text,artifact)
        # The branded hero uses a compact date label derived from the same dates.
        from hero import _date_label
        self.assertIn(_date_label(cfg['event']),hero)
        self.assertIn('data-start-date="'+cfg['event']['startDate']+'"',hero)
        self.assertIn('data-end-date="'+cfg['event']['endDate']+'"',hero)
        self.assertIn(facts['title'],hero);self.assertIn(facts['location'],hero)
        self.assertIn('20MB',posts+guide+doc)

    def test_resource_requires_exactly_one_destination(self):
        with self.assertRaises(ConfigError):self.run_generate(self.change('resources',{'C2':'https://example.org/two'}))
        with self.assertRaises(ConfigError):self.run_generate(self.change('resources',{'D2':'../outside.pdf'}))

    def test_late_registration_submissions_and_post_event_supplements_allowed(self):
        workbook=self.change('event',{'B15':'2027-03-19T17:00:00','B17':'2027-03-19T16:00:00','B18':'2027-04-02T23:59:00'})
        config,_=load_config(workbook)
        self.assertGreater(config['attendance']['closeAt'],config['event']['startAt'])
        self.assertGreater(config['submission']['closeAt'],config['event']['startAt'])
        self.assertGreater(config['submission']['supplementCloseAt'],config['event']['endAt'])
        self.run_generate(workbook)

    def test_optional_supplement_deadline_defaults_to_submission_close(self):
        config,_=load_config(self.change('event',{'B18':''}))
        self.assertEqual(config['submission']['supplementCloseAt'],config['submission']['closeAt'])
        self.run_generate(self.change('event',{'B18':''}))

    def test_v1_rejects_unsupported_english_interface(self):
        with self.assertRaises(ConfigError):
            load_config(self.change('event',{'B11':'en'}))

    def test_standalone_resource_link_and_canonical_workflow(self):
        config,_=load_config(self.workbook)
        resource=next(item for item in config['resources'] if item.get('asset'))
        self.assertEqual(resource_destination(config,resource),'https://conference.example.org'+resource['asset'])
        config['event']['siteUrl']=''
        self.assertEqual(resource_destination(config,resource),resource['asset'])
        self.assertIn('需补材料',next(item['answer'] for item in config['faqs'] if item['question']=='如何补充投稿附件？'))
        default,_=load_config(self.change('event',{'B19':''}))
        self.assertEqual(default['branding']['primaryColor'],'#5b9bd5')

    def test_symlink_output_is_rejected(self):
        actual=self.base/'actual';actual.mkdir();self.output.symlink_to(actual,target_is_directory=True)
        with self.assertRaises(ConfigError):self.run_generate()
        self.assertEqual(list(actual.iterdir()),[])


class V2HonestFactsTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(prefix='event-v2-facts-')
        self.base=Path(self.tmp.name)
        self.input=self.base/'input';shutil.copytree(ROOT/'input',self.input)
        self.workbook=self.input/'fictional-conference.xlsx'
        self.template=self.base/'template';self.template.mkdir()
        (self.template/'package.json').write_text('{}')
        self.output=self.base/'output'

    def tearDown(self):self.tmp.cleanup()

    def date_only(self):
        import openpyxl
        wb=openpyxl.load_workbook(self.workbook)
        for row in wb['event'].iter_rows(min_row=2):
            if row[0].value in ('event.startAt','event.endAt','event.capacity','attendance.openAt','attendance.closeAt','submission.openAt','submission.closeAt','submission.supplementCloseAt','home.target'):
                row[1].value=None
        wb['event'].append(['event.startDate',date(2027,3,18),'仅给出日期'])
        wb['event'].append(['event.endDate',date(2027,3,19),'仅给出日期'])
        wb['agenda']['B2']='全天';wb['agenda']['E2']=None
        path=self.input/'honest.xlsx';wb.save(path);return path

    def test_date_precision_unknown_capacity_and_windows_are_null(self):
        config,_=load_config(self.date_only())
        self.assertEqual(config['event']['startDate'],'2027-03-18')
        self.assertEqual(config['event']['endDate'],'2027-03-19')
        for key in ('startAt','endAt','capacity'):self.assertIsNone(config['event'][key])
        for group in ('attendance','submission'):
            self.assertIsNone(config[group]['openAt']);self.assertIsNone(config[group]['closeAt'])
        self.assertEqual(config['readiness']['mode'],'preview')
        self.assertEqual(len(config['readiness']['unresolved']),5)
        self.assertFalse(config['readiness']['attendanceEnabled']);self.assertFalse(config['readiness']['submissionEnabled'])
        self.assertEqual(config['agenda'][0]['time'],'全天')
        self.assertEqual(config['agenda'][0]['location'],'地点待通知')
        self.assertEqual(config['home']['target'],'')

    def test_preview_materials_omit_unknown_numeric_and_date_placeholders(self):
        generate(self.date_only(),self.output,self.template)
        with ZipFile(self.output/'materials/conference-guide.docx') as z:doc=z.read('word/document.xml').decode()
        texts=[doc,*[(self.output/'materials'/name).read_text() for name in ('poster.svg','social-posts.txt','conference-guide.html')]]
        for text in texts:
            for forbidden in ('1970','00:00-23:59','00:00—23:59','参会规模：1 人','参会规模：None','截止：None','适合参加：'):
                self.assertNotIn(forbidden,text)
            self.assertNotIn('会务预览',text)
            self.assertNotIn('报名与投稿待确认',text)
            self.assertNotIn('资料来源',text)
            self.assertIn('2027年03月18日至2027年03月19日',text)
        self.assertIn('全天',doc);self.assertIn('地点待通知',doc)

    def test_missing_time_is_semantic_and_known_bad_time_is_rejected(self):
        import openpyxl
        path=self.date_only();wb=openpyxl.load_workbook(path);wb['agenda']['B2']=None;wb.save(path)
        config,_=load_config(path);self.assertEqual(config['agenda'][0]['time'],'时间待通知')
        wb['agenda']['B2']='24:01';wb.save(path)
        with self.assertRaises(ConfigError):load_config(path)

    def test_independent_operation_readiness_and_date_conflict(self):
        import openpyxl
        path=self.date_only();wb=openpyxl.load_workbook(path)
        replacements={'event.capacity':42,'attendance.openAt':'2027-01-01T09:00:00','attendance.closeAt':'2027-03-17T18:00:00'}
        for row in wb['event'].iter_rows(min_row=2):
            if row[0].value in replacements:row[1].value=replacements[row[0].value]
        wb.save(path);config,_=load_config(path)
        self.assertTrue(config['readiness']['attendanceEnabled']);self.assertFalse(config['readiness']['submissionEnabled'])
        for row in wb['event'].iter_rows(min_row=2):
            if row[0].value=='event.startAt':row[1].value='2027-03-17T09:00:00'
        wb.save(path)
        with self.assertRaises(ConfigError):load_config(path)


    def test_html_direct_config_grouping_retains_all_distinct_speakers(self):
        from materials import agenda_html
        config,_=load_config(self.workbook)
        first,second=config['agenda'][:2]
        first['speakerGroup']=second['speakerGroup']='shared'
        first['speaker']='Source speaker A';second['speaker']='Source speaker B'
        html=agenda_html(config)
        self.assertIn('Source speaker A<br>Source speaker B',html)
        self.assertIn('rowspan="2"',html)

    def test_group_metadata_cannot_hide_different_or_nonadjacent_facts(self):
        import openpyxl
        for conflict in ('different_speaker','nonadjacent_speaker','different_location'):
            with self.subTest(conflict=conflict):
                wb=openpyxl.load_workbook(self.workbook)
                sheet=wb['agenda'];records=list(sheet.values)[1:];sheet.delete_rows(1,sheet.max_row)
                sheet.append(['date','time','title','speaker','chair','location','timeGroup','topicGroup','speakerGroup','chairGroup','locationGroup'])
                for index,record in enumerate(records):
                    values=[record[0],record[1],record[2],record[3] or '', '', record[4], '', '', '', '', '']
                    if conflict=='different_speaker' and index in (0,1):
                        values[3]='A' if index==0 else 'B';values[8]='shared-speaker'
                    elif conflict=='nonadjacent_speaker' and index in (0,2):
                        values[3]='A';values[8]='shared-speaker'
                    elif conflict=='different_location' and index in (0,1):
                        values[5]='A' if index==0 else 'B';values[10]='shared-venue'
                    sheet.append(values)
                path=self.input/'bad-group.xlsx';wb.save(path)
                with self.assertRaises(ConfigError):load_config(path)
                self.assertFalse(self.output.exists())

    @unittest.skipUnless((shutil.which('libreoffice') or shutil.which('soffice')) and Path('/home/agent/.codex/skills/builtins/documents/render_docx.py').is_file(),'Document renderer unavailable for optional rendered pagination regression')
    def test_rendered_34_row_long_guide_has_no_blank_or_horizontal_overflow(self):
        import subprocess
        import os
        import pdfplumber
        from generate import deterministic_docx
        config,_=load_config(self.workbook)
        # Synthetic long fixture avoids hiding a pre-extracted event inside test answers.
        from datetime import timedelta
        start=date.fromisoformat(config['event']['startDate'])
        config['event']['endDate']=(start+timedelta(days=4)).isoformat()
        config['event']['startAt']=config['event']['endAt']=None
        config['agenda']=[{'date':(start+timedelta(days=i//7)).isoformat(),'time':f'{9+i%7:02d}:00-{10+i%7:02d}:00','title':f'长议程测试主题 {i+1}','speaker':f'虚构讲者 {i+1}','chair':'虚构主持人','location':'虚构会场','chairGroup':f'day-{i//7}'} for i in range(34)]
        config['home']['intro'].append('长段落回归测试。'+('科研训练与学科交流的说明应当自然分页，不能将整节或整张表锁在同一页。'*24))
        path=self.base/'long-guide.docx';deterministic_docx(path,config)
        renderer=Path('/home/agent/.codex/skills/builtins/documents/render_docx.py')
        python=os.environ.get('CODEX_PRIMARY_RUNTIME_PYTHON',sys.executable)
        result=subprocess.run([python,str(renderer),str(path),'--output_dir',str(self.base/'render'),'--emit_pdf'],capture_output=True,text=True,timeout=180)
        self.assertEqual(result.returncode,0,result.stdout+result.stderr)
        pdf=next((self.base/'render').glob('*.pdf'))
        with pdfplumber.open(pdf) as document:
            self.assertGreater(len(document.pages),5)
            for index,page in enumerate(document.pages,1):
                text=page.extract_text() or ''
                self.assertGreater(len(text.strip()),35,f'Blank/sparse page {index}')
                for char in page.chars:
                    self.assertGreaterEqual(char['x0'],-1,f'Left overflow page {index}')
                    self.assertLessEqual(char['x1'],page.width+1,f'Right overflow page {index}')
                    self.assertGreaterEqual(char['top'],-1,f'Top overflow page {index}')
                    self.assertLessEqual(char['bottom'],page.height+1,f'Bottom overflow page {index}')
        self.assertTrue(list((self.base/'render').glob('page-*.png')))

if __name__=='__main__':unittest.main()
