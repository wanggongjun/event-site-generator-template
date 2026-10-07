"""Editable, source-faithful event materials. No operational facts are invented."""
from __future__ import annotations
import base64
from collections import OrderedDict
from datetime import date, datetime
import html
import io
import shutil
import subprocess
import tempfile
from pathlib import Path
from urllib.parse import quote
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED


def facts(config):
    from generate import event_facts
    return event_facts(config)


def resource_destination(config, item):
    if item.get('url'):
        return item['url']
    asset = item.get('asset', '')
    site = config['event'].get('siteUrl', '')
    return site.rstrip('/') + quote(asset, safe='/') if site else asset + '（本地演示路径；需先启动网站，不能作为独立公开链接）'


def public_fact_lines(config):
    f = facts(config)
    lines = [f['dates'], f['location']]
    if f['attendanceClose']:
        lines.append('参会报名截止：' + f['attendanceClose'])
    if f['submissionClose']:
        lines.append('学术投稿截止：' + f['submissionClose'])
    if f['capacity']:
        lines.append('参会规模：' + f['capacity'] + ' 人')
    if f['siteUrl']:
        lines.append('会议网站：' + f['siteUrl'])
    return lines


def draft_note(config):
    if config.get('readiness', {}).get('mode') != 'preview':
        return ''
    return '会务预览稿：报名或投稿条件尚未完整提供，待主办方确认后开放。'


def guide_sections(config):
    f = facts(config)
    overview = [f['dates'] + ' · ' + f['location'], config['travel']['address'], *config['home']['intro']]
    if config['home'].get('target'):
        overview.append('面向人群：' + config['home']['target'])
    overview.append('时区：' + config['event']['timezone'])
    if draft_note(config):
        overview.append(draft_note(config))
    organizations = []
    roles = OrderedDict()
    for item in config['home'].get('organizers', []):
        roles.setdefault(item['role'], []).append(item['name'])
    for role, names in roles.items():
        organizations.append(role + '：' + '、'.join(names))
    operational = []
    if f['capacity']:
        operational.append('参会规模：' + f['capacity'] + ' 人')
    for label, value in [('参会报名截止', f['attendanceClose']), ('学术投稿截止', f['submissionClose'])]:
        if value:
            operational.append(label + '：' + value)
    supplement = config['submission'].get('supplementCloseAt')
    if supplement:
        operational.append('补充材料截止：' + datetime.fromisoformat(supplement).strftime('%Y年%m月%d日 %H:%M'))
    operational += ['报名与投稿是两项独立流程。开放后，先完成个人信息，再分别提交。',
        '在线投稿系统每人一份投稿；最多 ' + str(config['files']['maxAttachments']) + ' 个附件，单个不超过 ' + str(config['files']['maxFileBytes']//1048576) + 'MB，支持 ' + '/'.join(config['files']['allowedExtensions']).upper() + '。']
    if draft_note(config):
        operational.insert(0, draft_note(config))
    if f['siteUrl']:
        operational.append('会议网站：' + f['siteUrl'])
    contacts = [item['name'] + '（' + item['responsibility'] + '）：' + ' / '.join(item[k] for k in ('email', 'phone') if item.get(k)) + (('。' + item['note']) if item.get('note') else '') for item in config['contacts']]
    meals = [item['title'] + '：' + item['body'] for item in config.get('mealGuide', [])]
    meals += [item['name'] + '：' + item['address'] + '。' + item['description'] + ((' ' + item['url']) if item.get('url') else '') for item in config['hotels']]
    sections = [
        ('会议简介', overview), ('组织机构', organizations),
        ('参会须知', [item['title'] + '：' + item['body'] for item in config.get('notices', [])]),
        ('会议日程', []), ('食宿指引', meals),
        ('交通指引', [config['travel']['address'], *[item['title'] + '：' + item['body'] for item in config['travel']['directions']]]),
        ('联系与常见问题', [*contacts, *[item['question'] + '\n' + item['answer'] for item in config['faqs']]]),
        ('报名与投稿系统说明', operational),
        ('园区地图与资料下载', [item['title'] + '：' + item['description'] + ' ' + resource_destination(config, item) for item in config['resources']]),
        ('资料来源', [item['title'] + '：' + item['body'] for item in config.get('sourceNotes', [])]),
    ]
    if config['event'].get('recordingsEnabled'):
        sections.append(('会议回放', [item['title'] + '：' + item['description'] + ' ' + item['url'] for item in config['recordings']]))
    return sections


def agenda_days(config):
    days = OrderedDict()
    for item in config['agenda']:
        days.setdefault(item['date'], []).append(item)
    return list(days.items())


def agenda_locations(items):
    groups = []
    for item in items:
        location = item.get('location') or '地点待通知'
        identity = item.get('locationGroup') or location
        if not groups or groups[-1][0] != identity:
            groups.append([identity, location, []])
        groups[-1][2].append(item)
    return groups


def agenda_html(config):
    escape = html.escape
    out = []
    for day, items in agenda_days(config):
        d = date.fromisoformat(day)
        rows = []
        for _, location, group in agenda_locations(items):
            rows.append('<tr class="venue"><th colspan="4">地点：' + escape(location) + '</th></tr>')
            spans, skipped = {}, set()
            for column, metadata in enumerate(('timeGroup','topicGroup','speakerGroup','chairGroup')):
                i=0
                while i<len(group):
                    identity=group[i].get(metadata);j=i+1
                    if identity:
                        while j<len(group) and group[j].get(metadata)==identity:j+=1
                    if identity and j-i>1 and all(group[k].get('speaker') or group[k].get('chair') for k in range(i,j)):
                        spans[(i,column)]=j-i
                        skipped.update((k,column) for k in range(i+1,j))
                    i=j
            for index,item in enumerate(group):
                values = [item.get('time') or '时间待通知', item['title'], item.get('speaker', ''), item.get('chair', '')]
                cells=[]
                merged_topic=not item.get('speaker') and not item.get('chair')
                for column,value in enumerate(values):
                    if (index,column) in skipped or (merged_topic and column in (2,3)):continue
                    attrs=(' colspan="3"' if merged_topic and column==1 else '')
                    if (index,column) in spans:
                        span=spans[(index,column)];attrs+=' rowspan="'+str(span)+'"'
                        field=('time','title','speaker','chair')[column]
                        # Defensive direct-config rendering retains all distinct
                        # source values, even outside the XLSX validator.
                        values_in_group=list(dict.fromkeys(group[k].get(field,'') for k in range(index,index+span)))
                        value='\n'.join(v for v in values_in_group if v)
                    cells.append('<td'+attrs+'>'+escape(value).replace('\n','<br>')+'</td>')
                rows.append('<tr>'+''.join(cells)+'</tr>')
        out.append('<div class="agenda-day"><h3>' + f'{d.year}年{d.month}月{d.day}日' + '</h3><table><colgroup><col style="width:19%"><col style="width:36%"><col style="width:28%"><col style="width:17%"></colgroup><thead><tr><th>时间</th><th>题目</th><th>主讲人</th><th>主持人</th></tr></thead><tbody>' + ''.join(rows) + '</tbody></table></div>')
    return ''.join(out)


def deterministic_docx(path, config, asset_root=None):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.7)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        style = doc.styles[name]
        for border in list(style._element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
        style.font.name = 'Noto Sans CJK SC'
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
        style.font.color.rgb = RGBColor(0,0,0)
    normal = doc.styles['Normal']
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.keep_with_next = False
    normal.paragraph_format.widow_control = True
    doc.styles['Title'].font.size = Pt(23)
    doc.styles['Title'].paragraph_format.space_after = Pt(9)
    for name, size in [('Heading 1', 15), ('Heading 2', 12)]:
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].paragraph_format.keep_with_next = True
        doc.styles[name].paragraph_format.space_before = Pt(12)
        doc.styles[name].paragraph_format.space_after = Pt(7)
    doc.add_heading(config['event']['title'], 0)
    if config['event'].get('subtitle'):
        doc.add_paragraph(config['event']['subtitle'])
    p = doc.add_paragraph('会议指南')
    p.paragraph_format.space_after = Pt(10)
    header = section.header.paragraphs[0]
    header.text = config['event']['shortTitle'] + ' 会议指南'
    header.runs[0].font.size = Pt(8)
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    run._r.addnext(field)
    def shade(cell, fill):
        el = OxmlElement('w:shd'); el.set(qn('w:fill'), fill); cell._tc.get_or_add_tcPr().append(el)
    def table_setup(table):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
        pr = table._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for edge in ('top','left','bottom','right','insideH','insideV'):
            el = OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'5');el.set(qn('w:color'),'D9D9D9');borders.append(el)
        pr.append(borders)
        for row in table.rows:
            # Only an individual row stays together. The entire table can flow.
            no_split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                margin = OxmlElement('w:tcMar')
                for edge, width in [('top',80),('bottom',80),('left',90),('right',90)]:
                    el=OxmlElement('w:'+edge);el.set(qn('w:w'),str(width));el.set(qn('w:type'),'dxa');margin.append(el)
                cell._tc.get_or_add_tcPr().append(margin)
                for paragraph in cell.paragraphs:
                    pf=paragraph.paragraph_format;pf.space_after=Pt(0);pf.line_spacing=1.1;pf.keep_with_next=False;pf.keep_together=False
                    for run in paragraph.runs:
                        run.font.size=Pt(9)
    def add_agenda():
        first = True
        for day, items in agenda_days(config):
            d=date.fromisoformat(day)
            heading = doc.add_heading(f'{d.year}年{d.month}月{d.day}日',2)
            # One continuous date heading, never a separate split date label.
            if not first and len(items) > 2:
                heading.paragraph_format.page_break_before = True
            first = False
            table = doc.add_table(rows=1, cols=4)
            widths = [1.27,2.42,1.89,1.29]
            for i, (label, width) in enumerate(zip(['时间','题目','主讲人','主持人'], widths)):
                table.columns[i].width = Inches(width)
                cell=table.rows[0].cells[i];cell.width=Inches(width);cell.text=label;shade(cell,'D9E5F2')
                cell.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in cell.paragraphs[0].runs:run.bold=True
            repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for _, location, group in agenda_locations(items):
                venue=table.add_row().cells; merged=venue[0].merge(venue[3]);merged.text='地点：'+location;shade(merged,'F2D4BE')
                merged.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
                merged.paragraphs[0].paragraph_format.keep_with_next=True
                for run in merged.paragraphs[0].runs:run.bold=True
                added_rows=[]
                for n,item in enumerate(group):
                    row=table.add_row();added_rows.append(row)
                    for i,value in enumerate([item.get('time') or '时间待通知',item['title'],item.get('speaker',''),item.get('chair','')]):
                        cell=row.cells[i];cell.width=Inches(widths[i]);cell.text=value
                        if i in (0,3):cell.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
                        if n%2:shade(cell,'F8FAFC')
                    if not item.get('speaker') and not item.get('chair'):
                        merged=row.cells[1].merge(row.cells[3]);merged.text=item['title']
                # Preserve the source's shared cells only when the caller
                # explicitly supplies the corresponding grouping metadata.
                for column, metadata in [(0,'timeGroup'),(1,'topicGroup'),(2,'speakerGroup'),(3,'chairGroup')]:
                    i=0
                    while i<len(group):
                        identity=group[i].get(metadata)
                        j=i+1
                        if identity:
                            while j<len(group) and group[j].get(metadata)==identity:j+=1
                        if identity and j-i>1 and all(len(added_rows[k]._tr.tc_lst)==4 for k in range(i,j)):
                            cell=added_rows[i].cells[column].merge(added_rows[j-1].cells[column])
                            value_field=['time','title','speaker','chair'][column]
                            values=list(dict.fromkeys(group[k].get(value_field,'') for k in range(i,j)))
                            cell.text='\n'.join(value for value in values if value)
                            if column in (0,3):cell.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
                        i=j
            table_setup(table)
            # Keep venue only with its next row; never set keep-next on all data.
            for row in table.rows[1:]:
                if len(row._tr.tc_lst)==1:
                    for paragraph in row.cells[0].paragraphs:paragraph.paragraph_format.keep_with_next=True
            # No trailing empty paragraph: it can spill to an otherwise blank
            # page immediately before the following page-break heading.
    for title, lines in guide_sections(config):
        if title=='会议日程':
            heading=doc.add_heading(title,1)
            heading.paragraph_format.page_break_before=True
            add_agenda()
            continue
        if not lines:
            continue
        heading=doc.add_heading(title,1)
        if title in ('参会须知','食宿指引'):heading.paragraph_format.page_break_before=True
        for line in lines:
            doc.add_paragraph(line)
        if title=='园区地图与资料下载' and asset_root:
            for item in config['resources']:
                if item.get('type')=='map' and item.get('asset'):
                    image_path=Path(asset_root)/item['asset'].removeprefix('/')
                    if image_path.is_file() and image_path.suffix.lower() in ('.png','.jpg','.jpeg'):
                        doc.add_picture(str(image_path),width=Inches(6.5))
    props=doc.core_properties
    props.title=config['event']['title']+'会议指南';props.author='';props.last_modified_by='';props.created=props.modified=datetime(2000,1,1)
    stream=io.BytesIO();doc.save(stream)
    with ZipFile(stream) as source, ZipFile(path,'w',ZIP_DEFLATED,compresslevel=9) as output:
        for name in sorted(source.namelist()):
            info=ZipInfo(name,date_time=(2000,1,1,0,0,0));info.compress_type=ZIP_DEFLATED;info.external_attr=0o600<<16
            output.writestr(info,source.read(name))


def _hero_data(hero_path):
    if not hero_path or not Path(hero_path).is_file():
        return ''
    path=Path(hero_path)
    media={'.webp':'image/webp','.jpg':'image/jpeg','.jpeg':'image/jpeg','.png':'image/png','.svg':'image/svg+xml'}.get(path.suffix.lower())
    return 'data:'+media+';base64,'+base64.b64encode(path.read_bytes()).decode('ascii') if media else ''


def poster_intro(config, max_units=38):
    paragraphs=config['home'].get('intro',[])
    if not paragraphs:return []
    text=paragraphs[1] if len(paragraphs)>1 else paragraphs[0]
    lines=[];line='';units=0
    for ch in text.replace('\n',' '):
        width=1 if ord(ch)>255 else .57
        if units+width>max_units and line:
            lines.append(line);line='';units=0
        line+=ch;units+=width
    if line:lines.append(line)
    return lines


def poster_svg(config, hero_path=None):
    escape=html.escape
    title=config['event']['title']
    lines=public_fact_lines(config)
    if draft_note(config):lines.append('会务预览稿，报名与投稿待确认')
    text=''.join(f'<text x="88" y="{490+i*78}" font-size="{min(32,900/max(1,sum(1 if ord(c)>255 else .57 for c in line))):.1f}">{escape(line)}</text>' for i,line in enumerate(lines))
    intro_y=490+len(lines)*78+58
    intro_lines=poster_intro(config)
    max_lines=max(1,int((1140-intro_y)/34))
    if len(intro_lines)>max_lines:
        intro_lines=intro_lines[:max_lines];intro_lines[-1]=intro_lines[-1][:-1]+'…'
    intro='<text x="88" y="'+str(intro_y-28)+'" font-size="27" font-weight="700">会议简介</text>'+''.join(f'<text x="88" y="{intro_y+i*34}" font-size="23">{escape(line)}</text>' for i,line in enumerate(intro_lines))
    hero=_hero_data(hero_path)
    image=f'<image x="0" y="0" width="1080" height="270" href="{hero}" preserveAspectRatio="xMidYMid meet"/>' if hero else ''
    return f'<svg xmlns="http://www.w3.org/2000/svg" width="1080" height="1440" viewBox="0 0 1080 1440"><title>{escape(title)}</title><rect width="1080" height="1440" fill="#071973"/>{image}<rect x="64" y="335" width="952" height="870" fill="#fff"/><g fill="#0e2e4f" font-family="Noto Sans CJK SC,Microsoft YaHei,sans-serif"><text x="88" y="405" font-size="40" font-weight="700">会议信息</text>{text}{intro}</g><g fill="white" font-family="Noto Sans CJK SC,Microsoft YaHei,sans-serif"><text x="88" y="1290" font-size="26">{escape(title)}</text><text x="88" y="1360" font-size="22">时间均以 {escape(config["event"]["timezone"])} 为准</text></g></svg>'


def poster_png(config, path, hero_path=None):
    from PIL import Image, ImageDraw, ImageFont
    fonts=[Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'),Path('/System/Library/Fonts/PingFang.ttc'),Path('C:/Windows/Fonts/msyh.ttc')]
    font_path=next((str(p) for p in fonts if p.is_file()),None)
    if not font_path:return False
    im=Image.new('RGB',(1080,1440),'#071973');draw=ImageDraw.Draw(im)
    if hero_path and Path(hero_path).is_file():
        source=Path(hero_path)
        if source.suffix.lower()=='.svg':
            try:
                import cairosvg
                data=cairosvg.svg2png(url=str(source),output_width=1080,output_height=270)
            except (ImportError, OSError):
                # Optional installed SVG renderer; never drop the brand hero.
                renderer=shutil.which('inkscape')
                if not renderer:return False
                with tempfile.TemporaryDirectory(prefix='event-hero-png-') as tmp:
                    output=Path(tmp)/'hero.png'
                    rendered=subprocess.run([renderer,str(source),'--export-type=png','--export-filename='+str(output),'--export-width=1080','--export-height=270'],capture_output=True,timeout=60)
                    if rendered.returncode or not output.is_file():return False
                    data=output.read_bytes()
            hero=Image.open(io.BytesIO(data)).convert('RGB')
        else:
            hero=Image.open(source).convert('RGB');hero.thumbnail((1080,270),Image.Resampling.LANCZOS)
        im.paste(hero,((1080-hero.width)//2,(270-hero.height)//2))
    draw.rectangle((64,335,1016,1205),fill='white')
    def write(text,y,size,color='#0e2e4f'):
        font=ImageFont.truetype(font_path,size)
        while draw.textbbox((0,0),text,font=font)[2]>900 and size>12:
            size-=1;font=ImageFont.truetype(font_path,size)
        draw.text((88,y),text,font=font,fill=color,anchor='lt')
    write('会议信息',360,40)
    lines=public_fact_lines(config)
    if draft_note(config):lines.append('会务预览稿，报名与投稿待确认')
    for i,line in enumerate(lines):write(line,450+i*78,32)
    intro_y=450+len(lines)*78+58
    intro_lines=poster_intro(config)
    max_lines=max(1,int((1100-intro_y)/34))
    if len(intro_lines)>max_lines:
        intro_lines=intro_lines[:max_lines];intro_lines[-1]=intro_lines[-1][:-1]+'…'
    write('会议简介',intro_y-35,27)
    for i,line in enumerate(intro_lines):write(line,intro_y+i*34,23)
    write(config['event']['title'],1250,26,'white');write('时间均以 '+config['event']['timezone']+' 为准',1320,22,'white')
    im.save(path,compress_level=9);return True


def write_materials(stage, config, hero_path=None):
    stage=Path(stage);folder=stage/'materials';folder.mkdir(parents=True,exist_ok=True)
    (folder/'poster.svg').write_text(poster_svg(config,hero_path),encoding='utf-8')
    png=poster_png(config,folder/'poster.png',hero_path)
    deterministic_docx(folder/'conference-guide.docx',config,stage/'frontend/public')
    sections=[]
    for title,lines in guide_sections(config):
        if title=='会议日程':content=agenda_html(config)
        else:content=''.join('<p>'+html.escape(line).replace('\n','<br>')+'</p>' for line in lines)
        if not content:continue
        if title=='园区地图与资料下载':
            for item in config['resources']:
                if item.get('type')=='map' and item.get('asset'):
                    content+='<img class="map" src="'+html.escape('../frontend/public'+item['asset'],quote=True)+'" alt="'+html.escape(item['title'],quote=True)+'">'
        sections.append('<section><h2>'+html.escape(title)+'</h2>'+content+'</section>')
    body='<!doctype html><html lang="'+config['event']['language']+'"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'+html.escape(config['event']['title'])+'会议指南</title><style>body{max-width:900px;margin:40px auto;padding:0 24px;color:#193254;font:16px/1.7 system-ui,sans-serif}h1,h2,h3{color:#000;line-height:1.4}h2,h3{break-after:avoid}p{overflow-wrap:anywhere}table{border-collapse:collapse;width:100%;table-layout:fixed;margin:12px 0 28px}th,td{border:1px solid #d9d9d9;padding:8px;vertical-align:middle;overflow-wrap:anywhere}thead{display:table-header-group;background:#d9e5f2}tr{break-inside:avoid}td:first-child{text-align:center}.venue th{background:#f2d4be}.map{max-width:100%}@media print{body{margin:0;font-size:10.5pt}.agenda-day{break-before:page}.agenda-day:first-child{break-before:auto}}</style><body><h1>'+html.escape(config['event']['title'])+'</h1><p>会议指南</p>'+''.join(sections)+'</body></html>'
    (folder/'conference-guide.html').write_text(body,encoding='utf-8')
    f=facts(config);long=[f['title']]
    if f['subtitle']:long.append(f['subtitle'])
    long += ['时间：'+f['dates'],'地点：'+f['location'],*config['home']['intro']]
    if config['home'].get('target'):long.append('面向人群：'+config['home']['target'])
    long += public_fact_lines(config)[2:]
    if draft_note(config):long.append(draft_note(config))
    long.append('所有时间以 '+config['event']['timezone']+' 为准。')
    short=f['title']+'将于'+f['dates']+'在'+f['location']+'举行。'
    if f['siteUrl']:short+='详情：'+f['siteUrl']
    if draft_note(config):short+=' '+draft_note(config)
    posts='长版宣传文案'+(' 草稿' if draft_note(config) else '')+'\n'+'\n'.join(long)+'\n\n简版宣传文案'+(' 草稿' if draft_note(config) else '')+'\n'+short+'\n'
    (folder/'social-posts.txt').write_text(posts,encoding='utf-8');return png
