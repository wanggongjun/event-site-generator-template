#!/usr/bin/env python3
"""Deterministic canonical XLSX -> portable event app and editable materials.

Only workbook cells and referenced local public assets are event-specific inputs.
The runtime template is copied unchanged; existing business state is never deleted.
"""
from __future__ import annotations
import argparse
import hashlib
import html
import io
import json
import math
import os
from pathlib import Path, PurePosixPath
import re
import shutil
import sys
import tempfile
from datetime import date, datetime, time, timezone
from urllib.parse import urlsplit, quote
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError
from zipfile import ZipFile, ZIP_DEFLATED
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
VERSION = '4.0.0'
MAX_WORKBOOK_BYTES = 10 * 1024 * 1024
MAX_ASSET_BYTES = 20 * 1024 * 1024
SAFE_ASSET_EXTENSIONS = {'.png', '.jpg', '.jpeg', '.webp', '.pdf', '.docx', '.pptx'}
FORBIDDEN_SEGMENTS = {'.git', 'node_modules', 'vendor', 'data', 'uploads', '__pycache__', 'test-evidence', 'offline-preview', '.aws', '.codex', '.agents'}
SECRET_KEY = re.compile(r'(password|passwd|secret|token|credential|api.?key|private.?key|authorization|database.?url|smtp|webhook|feishu|oauth)', re.I)
EVENT_FIELDS = {
 'event.slug': ('str', True), 'event.title': ('str', True), 'event.shortTitle': ('str', True),
 'event.subtitle': ('str', False), 'event.startDate': ('date', False), 'event.endDate': ('date', False),
 'event.startAt': ('datetime', False), 'event.endAt': ('datetime', False),
 'event.timezone': ('str', True), 'event.location': ('str', True), 'event.capacity': ('positive-int', False),
 'event.language': ('str', False), 'event.siteUrl': ('url', False), 'event.recordingsEnabled': ('bool', False),
 'attendance.openAt': ('datetime', False), 'attendance.closeAt': ('datetime', False),
 'submission.openAt': ('datetime', False), 'submission.closeAt': ('datetime', False), 'submission.supplementCloseAt': ('datetime', False),
 'branding.primaryColor': ('color', False), 'branding.seriesText': ('str', False), 'branding.bannerTitle': ('str', False), 'branding.heroImage': ('asset', False), 'branding.logo': ('asset', False),
 'home.target': ('str', False), 'travel.address': ('str', True), 'travel.mapUrl': ('url', False),
}
COLLECTIONS = {
 'intro': (['text'], ['text']),
 'brandmarks': (['label', 'asset'], ['label', 'asset']),
 'features': (['title', 'description'], ['title', 'description']),
 'organizers': (['role', 'name'], ['role', 'name']),
 'agenda': (['date', 'time', 'title', 'speaker', 'chair', 'location', 'timeGroup', 'topicGroup', 'speakerGroup', 'chairGroup', 'locationGroup'], ['date', 'title']),
 'notices': (['title', 'body'], ['title', 'body']),
 'mealGuide': (['title', 'body'], ['title', 'body']),
 'sourceNotes': (['title', 'body'], ['title', 'body']),
 'directions': (['title', 'body'], ['title', 'body']),
 'hotels': (['name', 'address', 'description', 'url'], ['name', 'address', 'description']),
 'contacts': (['name', 'responsibility', 'email', 'phone', 'note'], ['name', 'responsibility']),
 'faqs': (['question', 'answer'], ['question', 'answer']),
 'resources': (['title', 'description', 'url', 'asset', 'type'], ['title', 'description']),
 'recordings': (['title', 'description', 'url'], ['title', 'description', 'url']),
}
DEFAULT_FIELDS = {'profileOptionalFields', 'files.maxFileBytes', 'files.maxAttachments', 'files.allowedExtensions', 'sync.pollSeconds'}
DEFAULTS = {'profileOptionalFields': ['department', 'job', 'personalIntroduction'], 'files.maxFileBytes': 20971520,
            'files.maxAttachments': 3, 'files.allowedExtensions': ['pdf', 'docx', 'pptx'], 'sync.pollSeconds': 60}

class ConfigError(ValueError):
    pass

def fail(message):
    raise ConfigError(message)

def digest(data):
    return hashlib.sha256(data).hexdigest()

def plain(value, label, allow_empty=False):
    if value is None or value == '':
        if allow_empty:
            return ''
        fail(f'{label}: 必填内容为空')
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return 'true' if value else 'false'
    if not isinstance(value, (str, int, float)):
        fail(f'{label}: 需要文本、数字或日期')
    text = str(value).strip()
    if len(text) > 10000 or any(ord(ch) < 32 and ch not in '\n\t' for ch in text):
        fail(f'{label}: 文本过长或包含控制字符')
    if not text and not allow_empty:
        fail(f'{label}: 必填内容为空')
    return text

def parse_bool(value, label):
    if isinstance(value, bool):
        return value
    v = plain(value, label).lower()
    if v not in ('true', 'false', '1', '0', '是', '否'):
        fail(f'{label}: 布尔值必须为 true/false、1/0 或 是/否')
    return v in ('true', '1', '是')

def positive_int(value, label):
    if isinstance(value, bool):
        fail(f'{label}: 需要正整数')
    v = plain(value, label)
    if not re.fullmatch(r'\d+(?:\.0+)?', v):
        fail(f'{label}: 需要正整数')
    number = int(float(v))
    if number <= 0 or number > 2147483647:
        fail(f'{label}: 正整数超出范围')
    return number

def safe_url(value, label):
    v = plain(value, label)
    try:
        u = urlsplit(v)
        _ = u.port
    except ValueError:
        fail(f'{label}: URL 格式无效')
    if u.scheme not in ('http', 'https') or not u.hostname or u.username or u.password or '\\' in v or any(ch.isspace() for ch in v):
        fail(f'{label}: 只允许无凭据的完整 http/https URL')
    return v

def parse_datetime(value, label, zone):
    v = plain(value, label)
    try:
        dt = datetime.fromisoformat(v.replace('Z', '+00:00'))
    except ValueError:
        fail(f'{label}: 需要 ISO 日期时间，例如 2027-03-18T09:00:00')
    if 'T' not in v and ' ' not in v:
        fail(f'{label}: 必须包含日期及时间')
    if dt.tzinfo is None:
        candidate = dt.replace(tzinfo=zone)
        if candidate.astimezone(timezone.utc).astimezone(zone).replace(tzinfo=None) != dt:
            fail(f'{label}: 该本地时间因夏令时不存在')
        if dt.replace(tzinfo=zone, fold=0).utcoffset() != dt.replace(tzinfo=zone, fold=1).utcoffset():
            fail(f'{label}: 该本地时间有夏令时歧义，请提供明确 UTC 偏移')
        dt = candidate
    else:
        if dt.utcoffset() != dt.astimezone(zone).utcoffset():
            fail(f'{label}: UTC 偏移与 event.timezone 不一致')
        dt = dt.astimezone(zone)
    return dt.isoformat(timespec='seconds')

def parse_date(value, label):
    if isinstance(value, datetime):
        if value.time() != time(0):
            fail(f'{label}: 日期字段不接受带具体时刻的值，请另填 startAt/endAt')
        value = value.date()
    try:
        return date.fromisoformat(plain(value, label)).isoformat()
    except ValueError:
        fail(f'{label}: 需要 YYYY-MM-DD 日期')

def readiness(config):
    """Configuration readiness is independent of the current clock and sign-in."""
    unresolved = []
    for group, field in [('event', 'capacity'), ('attendance', 'openAt'), ('attendance', 'closeAt'), ('submission', 'openAt'), ('submission', 'closeAt')]:
        if config[group].get(field) is None:
            unresolved.append({'key': f'{group}.{field}', 'reason': '未提供会务事实，需主办方确认'})
    capacity = config['event'].get('capacity') is not None
    enabled = lambda group: capacity and all(config[group].get(field) is not None for field in ('openAt', 'closeAt'))
    return {'mode': 'preview' if unresolved else 'ready', 'unresolved': unresolved,
            'attendanceEnabled': enabled('attendance'), 'submissionEnabled': enabled('submission')}

def safe_asset(workbook_dir, value, label):
    v = plain(value, label)
    p = PurePosixPath(v)
    if '\\' in v or p.is_absolute() or any(part in ('..', '.', '') or part.startswith('.') for part in p.parts) or ':' in v:
        fail(f'{label}: 素材必须是工作簿目录内的安全相对路径')
    if not p.parts or any(part in FORBIDDEN_SEGMENTS for part in p.parts) or p.suffix.lower() not in SAFE_ASSET_EXTENSIONS:
        fail(f'{label}: 只允许 png/jpg/jpeg/webp/pdf/docx/pptx 公共素材')
    src = workbook_dir.joinpath(*p.parts)
    base = workbook_dir.resolve()
    if any(piece.is_symlink() for piece in [src, *list(src.parents)[:len(p.parts) - 1]]):
        fail(f'{label}: 禁止链接素材')
    try:
        src.resolve().relative_to(base)
    except ValueError:
        fail(f'{label}: 素材越出工作簿目录')
    if not src.is_file() or src.stat().st_size > MAX_ASSET_BYTES:
        fail(f'{label}: 素材不存在或超过 20MB')
    data = src.read_bytes()
    if p.suffix.lower() in ('.png', '.jpg', '.jpeg', '.webp'):
        from PIL import Image
        try:
            with Image.open(io.BytesIO(data)) as im:
                if im.width * im.height > 40000000:
                    fail(f'{label}: 图片像素过多')
                im.verify()
        except ConfigError:
            raise
        except Exception:
            fail(f'{label}: 图片文件无效')
    elif p.suffix.lower() == '.pdf':
        if not data.startswith(b'%PDF-'):
            fail(f'{label}: PDF 文件头无效')
    else:
        try:
            with ZipFile(io.BytesIO(data)) as z:
                if any('vba' in f.filename.lower() or f.filename.startswith('/') or '..' in PurePosixPath(f.filename).parts for f in z.infolist()):
                    fail(f'{label}: 禁止宏或不安全文档内容')
                if '[Content_Types].xml' not in z.namelist():
                    fail(f'{label}: Office 文件无效')
        except ConfigError:
            raise
        except Exception:
            fail(f'{label}: Office 文件无效')
    return '/assets/' + p.as_posix(), data

def read_rows(sheet, expected):
    if sheet.max_row > 2000 or sheet.max_column > 30:
        fail(f'{sheet.title}: 工作表过大')
    rows = list(sheet.iter_rows(values_only=False))
    if not rows:
        fail(f'{sheet.title}: 缺少表头')
    header = [plain(c.value, f'{sheet.title}!{c.coordinate}', True) for c in rows[0]]
    while header and not header[-1]:
        header.pop()
    if header != expected:
        fail(f'{sheet.title}: 第 1 行表头必须精确为 {",".join(expected)}')
    output = []
    for row in rows[1:]:
        if not any(c.value not in (None, '') for c in row):
            continue
        if any(c.data_type == 'f' or c.data_type == 'e' for c in row):
            fail(f'{sheet.title}!{row[0].row}: 配置不允许公式或错误单元格，请写入最终值')
        if any(c.value not in (None, '') for c in row[len(expected):]):
            fail(f'{sheet.title}!{row[0].row}: 存在未定义列的数据')
        output.append((row[0].row, {name: row[i].value if i < len(row) else None for i, name in enumerate(expected)}))
    return output

def load_config(workbook_path):
    """Validate everything without touching an output directory. Return config/assets."""
    path = Path(workbook_path).resolve()
    if path.suffix.lower() != '.xlsx' or not path.is_file() or path.stat().st_size > MAX_WORKBOOK_BYTES:
        fail('请输入存在且不超过 10MB 的 .xlsx 工作簿')
    raw = path.read_bytes()
    try:
        with ZipFile(io.BytesIO(raw)) as z:
            if sum(i.file_size for i in z.infolist()) > 50 * 1024 * 1024 or any('vbaProject' in i.filename for i in z.infolist()):
                fail('工作簿包含宏或解压体积超过 50MB')
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=False, data_only=False, keep_links=False)
    except ConfigError:
        raise
    except Exception as exc:
        fail(f'无法读取 XLSX: {type(exc).__name__}')
    allowed = {'event', 'defaults', *COLLECTIONS}
    if set(wb.sheetnames) - allowed:
        fail('未知工作表: ' + ', '.join(sorted(set(wb.sheetnames) - allowed)))
    if 'event' not in wb.sheetnames:
        fail('缺少 event 工作表')
    values = {}
    for row, record in read_rows(wb['event'], ['key', 'value', 'description']):
        key = plain(record['key'], f'event!A{row}')
        if SECRET_KEY.search(key):
            fail(f'event!A{row}: 禁止在配置内存储凭据或集成密钥')
        if key not in EVENT_FIELDS:
            fail(f'event!A{row}: 未定义配置项 {key}')
        if key in values:
            fail(f'event!A{row}: 重复配置项 {key}')
        values[key] = record['value']
    for key, (_, required) in EVENT_FIELDS.items():
        if required and values.get(key) in (None, ''):
            fail(f'event: 缺少必填配置 {key}')
    timezone_name = plain(values['event.timezone'], 'event.timezone')
    try:
        zone = ZoneInfo(timezone_name)
    except (ZoneInfoNotFoundError, ValueError):
        fail('event.timezone: 无效 IANA 时区（例如 Asia/Shanghai）')
    config = {'event': {}, 'attendance': {}, 'submission': {}, 'branding': {}, 'home': {}, 'travel': {}}
    assets = {}
    for key, (kind, _) in EVENT_FIELDS.items():
        val = values.get(key)
        if val in (None, ''):
            continue
        if kind == 'datetime':
            val = parse_datetime(val, key, zone)
        elif kind == 'date':
            val = parse_date(val, key)
        elif kind == 'positive-int':
            val = positive_int(val, key)
        elif kind == 'bool':
            val = parse_bool(val, key)
        elif kind == 'url':
            val = safe_url(val, key)
        elif kind == 'asset':
            if Path(str(val)).suffix.lower() not in ('.png','.jpg','.jpeg','.webp'):
                fail(f'{key}: 品牌素材需要 png/jpg/jpeg/webp 图片')
            val, data = safe_asset(path.parent, val, key)
            assets[val.removeprefix('/')] = data
        else:
            val = plain(val, key)
            if kind == 'color' and not re.fullmatch(r'#[0-9a-fA-F]{6}', val):
                fail(f'{key}: 必须为 #RRGGBB 颜色')
        group, name = key.split('.')
        config[group][name] = val
    e = config['event']
    if not re.fullmatch(r'[a-z0-9]+(?:-[a-z0-9]+)*', e['slug']) or len(e['slug']) > 64:
        fail('event.slug: 需要不超过 64 字符的小写字母、数字和单连字符')
    e.setdefault('subtitle', '')
    e.setdefault('language', 'zh-CN')
    if e['language'] != 'zh-CN':
        fail('event.language: 当前仅支持 zh-CN 中文界面')
    e.setdefault('siteUrl', '')
    e.setdefault('recordingsEnabled', False)
    config['branding'].setdefault('primaryColor', '#5b9bd5')
    config['branding'].setdefault('heroImage', '/assets/hero.svg')
    for field in ('startAt', 'endAt', 'capacity'):
        e.setdefault(field, None)
    for precise, day in [('startAt', 'startDate'), ('endAt', 'endDate')]:
        if e.get(precise):
            derived = datetime.fromisoformat(e[precise]).date().isoformat()
            if e.get(day) and e[day] != derived:
                fail(f'event: {day} 与 {precise} 日期不一致')
            e[day] = derived
        elif not e.get(day):
            fail(f'event: 需要 {day} 或含具体时刻的 {precise}')
    start_date, end_date = (date.fromisoformat(e[k]) for k in ('startDate', 'endDate'))
    if start_date > end_date:
        fail('event: startDate 不可晚于 endDate')
    if e['startAt'] and e['endAt'] and datetime.fromisoformat(e['startAt']) >= datetime.fromisoformat(e['endAt']):
        fail('event: startAt 必须早于 endAt')
    for group in ('attendance', 'submission'):
        for field in ('openAt', 'closeAt'):
            config[group].setdefault(field, None)
        if config[group]['openAt'] and config[group]['closeAt']:
            opened, closed = (datetime.fromisoformat(config[group][k]) for k in ('openAt', 'closeAt'))
            if opened >= closed:
                fail(f'{group}: 需要 openAt < closeAt')
    config['submission'].setdefault('supplementCloseAt', config['submission']['closeAt'])
    if config['submission']['supplementCloseAt']:
        if not config['submission']['closeAt']:
            fail('submission: 填写 supplementCloseAt 时必须提供 closeAt')
        if datetime.fromisoformat(config['submission']['supplementCloseAt']) < datetime.fromisoformat(config['submission']['closeAt']):
            fail('submission: 需要 closeAt <= supplementCloseAt')
    config['home'].setdefault('target', '')
    config['readiness'] = readiness(config)
    defaults = dict(DEFAULTS)
    if 'defaults' in wb.sheetnames:
        seen = set()
        for row, record in read_rows(wb['defaults'], ['key', 'value', 'description']):
            key = plain(record['key'], f'defaults!A{row}')
            if SECRET_KEY.search(key):
                fail('defaults: 禁止凭据配置')
            if key not in DEFAULT_FIELDS or key in seen:
                fail(f'defaults!A{row}: 未定义或重复配置项 {key}')
            seen.add(key)
            value = record['value']
            if value in (None, ''):
                continue
            if key in ('profileOptionalFields', 'files.allowedExtensions'):
                items = [item.strip() for item in plain(value, key).split(',') if item.strip()]
                allowed_items = set(DEFAULTS[key])
                if len(items) != len(set(items)) or set(items) - allowed_items:
                    fail(f'{key}: 重复或不支持的选项')
                if key == 'files.allowedExtensions' and not items:
                    fail(f'{key}: 不可为空')
                defaults[key] = items
            else:
                defaults[key] = positive_int(value, key)
    if defaults['files.maxFileBytes'] > 20971520 or defaults['files.maxAttachments'] > 3 or not 5 <= defaults['sync.pollSeconds'] <= 3600:
        fail('defaults: 附件上限不可超过 20MB/3 个；轮询间隔必须为 5–3600 秒')
    config['files'] = {key.split('.')[1]: defaults[key] for key in DEFAULT_FIELDS if key.startswith('files.')}
    config['sync'] = {'pollSeconds': defaults['sync.pollSeconds']}
    config['forms'] = {'profileOptionalFields': defaults['profileOptionalFields']}
    for name, (columns, required_columns) in COLLECTIONS.items():
        result, identities = [], set()
        if name in wb.sheetnames:
            if name == 'agenda':
                header = [c.value for c in wb[name][1]]
                while header and header[-1] is None:
                    header.pop()
                if header == ['date', 'time', 'title', 'speaker', 'location']:
                    columns = header  # v1 workbooks remain valid without invented chair data.
                elif header == ['date', 'time', 'title', 'speaker', 'chair', 'location']:
                    columns = header
            for row, record in read_rows(wb[name], columns):
                item = {}
                for field in columns:
                    v = plain(record[field], f'{name}!{field}{row}', field not in required_columns)
                    if not v:
                        continue
                    if field == 'url':
                        v = safe_url(v, f'{name}!{field}{row}')
                    elif field == 'asset':
                        if name == 'brandmarks' and Path(v).suffix.lower() not in ('.png','.jpg','.jpeg','.webp'):
                            fail(f'brandmarks!{row}: 品牌标识需要图片素材')
                        v, data = safe_asset(path.parent, v, f'{name}!{field}{row}')
                        if v in ('/assets/hero.svg',):
                            fail('素材路径与生成文件冲突')
                        assets[v.removeprefix('/')] = data
                    item[field] = v
                if name == 'resources' and bool(item.get('url')) == bool(item.get('asset')):
                    fail(f'resources!{row}: url 和 asset 必须且只能填写一项')
                if name == 'contacts':
                    if not item.get('email') and not item.get('phone'):
                        fail(f'contacts!{row}: email 或 phone 至少填写一项')
                    if item.get('email') and not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', item['email']):
                        fail(f'contacts!{row}: email 格式无效')
                if name == 'agenda':
                    item['date'] = parse_date(record['date'], f'agenda!date{row}')
                    item.setdefault('time', '时间待通知')
                    item.setdefault('location', '地点待通知')
                    item.setdefault('speaker', '')
                    item.setdefault('chair', '')
                    if item['time'] not in ('全天', '时间待通知'):
                        try:
                            if not re.fullmatch(r'\d{2}:\d{2}(?:-\d{2}:\d{2})?', item['time']):
                                raise ValueError()
                            times = item['time'].split('-')
                            for t in times:
                                time.fromisoformat(t)
                            if len(times) == 2 and times[0] >= times[1]:
                                raise ValueError()
                        except ValueError:
                            fail(f'agenda!{row}: time 需要 全天、时间待通知、HH:MM 或有序 HH:MM-HH:MM')
                    if not start_date <= date.fromisoformat(item['date']) <= end_date:
                        fail(f'agenda!{row}: 日程日期超出活动日期')
                    identity = (item['date'], item['time'], item['location'], item['title'])
                elif name == 'organizers':
                    identity = (item['role'], item['name'])
                elif name == 'contacts':
                    identity = (item['name'], item['responsibility'])
                else:
                    identity = item[columns[0]]
                if identity in identities:
                    fail(f'{name}!{row}: 重复或冲突记录')
                identities.add(identity)
                result.append(item)
        if name == 'intro':
            config['home']['intro'] = [item['text'] for item in result]
        elif name in ('features', 'organizers'):
            config['home'][name] = result
        elif name == 'directions':
            config['travel']['directions'] = result
        else:
            config[name] = result
    # Group IDs preserve source merged cells; they must never hide a different
    # fact or join non-adjacent rows in the same date/location block.
    for field, metadata in [('time','timeGroup'),('title','topicGroup'),('speaker','speakerGroup'),('chair','chairGroup'),('location','locationGroup')]:
        groups = {}
        for index, item in enumerate(config['agenda']):
            identity = item.get(metadata)
            if not identity:
                continue
            key = (item['date'], identity)
            if key in groups:
                previous_value, previous_index = groups[key]
                if previous_value != item.get(field, '') or previous_index != index-1:
                    fail(f'agenda: {metadata} 必须表示同日相邻且内容一致的 {field}')
            groups[key] = (item.get(field, ''), index)
    if len(config.get('brandmarks', [])) > 5:
        fail('brandmarks: 固定横幅品牌区最多接收五个标识')
    if not e['recordingsEnabled']:
        config['recordings'] = []
    if not config['home']['intro'] or not config['agenda'] or not config['contacts'] or not config['faqs']:
        fail('intro、agenda、contacts、faqs 各至少需要一条记录')
    from text_refs import resolve_content_refs
    from hero_binding import apply_hero_binding
    try:
        resolve_content_refs(config)
    except ValueError as exc:
        fail(str(exc))
    apply_hero_binding(config, assets, path)
    config['source'] = {'schemaVersion': VERSION, 'workbookSha256': digest(raw)}
    wb.close()
    return config, assets

def event_facts(config):
    e = config['event']
    start, end = (date.fromisoformat(e[k]) for k in ('startDate', 'endDate'))
    date_text = f'{start:%Y年%m月%d日}' if start == end else f'{start:%Y年%m月%d日}至{end:%Y年%m月%d日}'
    close_text = lambda group: datetime.fromisoformat(config[group]['closeAt']).strftime('%Y年%m月%d日 %H:%M') if config[group].get('closeAt') else None
    return {'title': e['title'], 'subtitle': e.get('subtitle', ''), 'dates': date_text,
            'location': e['location'], 'capacity': str(e['capacity']) if e.get('capacity') else None,
            'siteUrl': e.get('siteUrl', ''), 'attendanceClose': close_text('attendance'),
            'submissionClose': close_text('submission')}

def svg_text(value):
    return html.escape(str(value), quote=True)

def text_width(text):
    return sum(1.0 if ord(c) > 255 else 0.57 for c in text)

def hero_svg(config, assets=None):
    from hero import hero_svg as branded_hero_svg
    return branded_hero_svg(config, assets)

# Material functions live separately to keep source-aware pagination independent
# from validation and safe file-overlay mechanics. Public imports remain stable.
from materials import poster_svg, poster_png, guide_sections, deterministic_docx, write_materials, resource_destination

def template_files(template):
    if not template.is_dir():
        fail('缺少运行时 template/ 目录')
    result=[]
    for p in sorted(template.rglob('*')):
        rel=p.relative_to(template)
        if any(part in FORBIDDEN_SEGMENTS or part == 'dist' or (part.startswith('.') and part not in ('.env.example','.gitignore','.npmrc')) for part in rel.parts):
            continue
        if p.is_symlink():
            fail(f'template: 禁止符号链接 {rel}')
        if p.is_file():
            result.append((rel,p.read_bytes()))
    if not result:
        fail('template/ 为空')
    return result

def existing_output(output,slug):
    if output.is_symlink():
        fail('输出目录不可为符号链接')
    if not output.exists():
        return None
    if not output.is_dir():
        fail('输出路径已存在且不是目录')
    if not any(output.iterdir()):
        return None
    manifest_path=output/'manifest.json'
    if not manifest_path.is_file() or manifest_path.is_symlink():
        fail('拒绝写入非空未知目录；请使用新的空目录或本生成器已有输出')
    try:
        manifest=json.loads(manifest_path.read_text(encoding='utf-8'))
    except (ValueError,OSError):
        fail('已有 manifest.json 无效')
    if manifest.get('generator') != 'canonical-xlsx-event-template' or manifest.get('eventSlug') != slug:
        fail('已有输出不属于本生成器，或 event.slug 已改变；不同会议必须使用独立输出目录')
    return manifest

def public_config(config):
    """The public product contains event content, never generation/connection diagnostics."""
    result = {key: value for key, value in config.items() if key not in {'source', 'sourceNotes', 'readiness', 'sync'}}
    result['branding'] = {key: value for key, value in config.get('branding', {}).items() if key not in {'heroWarning', 'heroBinding', 'heroSourceImage'}}
    return result

def generate(workbook_path, output=None, template=None):
    config, assets=load_config(workbook_path)
    slug=config['event']['slug']; target=Path(output) if output else ROOT/'generated'/slug
    target=target.absolute();template=Path(template) if template else ROOT/'template'
    for parent in [target,*target.parents]:
        if parent.is_symlink():
            fail('输出目录的父路径不可为符号链接')
    try:
        target.resolve().relative_to(template.resolve())
        fail('输出目录不可位于 template/ 内')
    except ValueError:
        pass
    old=existing_output(target,slug)
    files=template_files(template)
    for rel,_ in files:
        if any(part in ('data','uploads') for part in rel.parts) or rel.name == '.env':
            fail('运行时模板包含业务状态或凭据')
    # All input/config/template validation has finished before output is touched.
    with tempfile.TemporaryDirectory(prefix='event-stage-') as tmp:
        stage=Path(tmp)
        for rel,data in files:
            p=stage/rel;p.parent.mkdir(parents=True,exist_ok=True);p.write_bytes(data)
        config_bytes=(json.dumps(config,ensure_ascii=False,indent=2,sort_keys=True)+'\n').encode('utf-8')
        (stage/'config.json').write_bytes(config_bytes)
        public_bytes=(json.dumps(public_config(config),ensure_ascii=False,indent=2,sort_keys=True)+'\n').encode('utf-8')
        for rel in ('frontend/public/config.json','frontend/src/config.generated.json'):
            p=stage/rel;p.parent.mkdir(parents=True,exist_ok=True);p.write_bytes(public_bytes)
        assets['assets/hero.svg']=hero_svg(config,assets).encode('utf-8')
        for rel,data in assets.items():
            p=stage/'frontend/public'/rel;p.parent.mkdir(parents=True,exist_ok=True);p.write_bytes(data)
        hero_path=stage/'frontend/public'/config['branding']['heroImage'].removeprefix('/')
        png=write_materials(stage,config,hero_path=hero_path)
        artifact_paths=[p.relative_to(stage).as_posix() for p in sorted(stage.rglob('*')) if p.is_file()]
        # Existing arbitrary user files must not be overwritten accidentally.
        previous_artifacts=set(old.get('artifacts',[])) if old else set()
        for rel in artifact_paths + ['manifest.json']:
            dest=target/rel
            if any(p.is_symlink() for p in [dest,*list(dest.parents)[:len(PurePosixPath(rel).parts)]]):
                fail(f'拒绝写入符号链接: {rel}')
            if dest.exists() and not dest.is_file():
                fail(f'输出文件路径被目录占用: {rel}')
            if old and dest.exists() and rel not in previous_artifacts and rel != 'manifest.json':
                fail(f'拒绝覆盖非生成文件: {rel}')
        manifest={'generator':'canonical-xlsx-event-template','generatorVersion':VERSION,'eventSlug':slug,
                  'workbookSha256':config['source']['workbookSha256'],
                  'configSha256':digest(config_bytes),'publicConfigSha256':digest(public_bytes),'templateSha256':digest(b''.join(str(r).encode()+b'\0'+d for r,d in files)),
                  'artifacts':artifact_paths,'sha256':{rel:digest((stage/rel).read_bytes()) for rel in artifact_paths},
                  'readiness':config['readiness'],
                  'warnings':([config['branding']['heroWarning']] if config['branding'].get('heroWarning') else []) + ([('会务预览：缺少 ' + '、'.join(item['key'] for item in config['readiness']['unresolved']))] if config['readiness']['mode']=='preview' else []) + ([] if png else ['PNG 未生成：系统缺少中文字体或 SVG 栅格化工具；可编辑 SVG 已生成。'])}
        (stage/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2,sort_keys=True)+'\n',encoding='utf-8')
        target.mkdir(parents=True,exist_ok=True)
        # Overlay only generator-owned files. No delete/rmtree occurs in target.
        for rel in artifact_paths + ['manifest.json']:
            src=stage/rel;dest=target/rel;dest.parent.mkdir(parents=True,exist_ok=True)
            fd,tmp_name=tempfile.mkstemp(prefix='.event-write-',dir=dest.parent)
            try:
                with os.fdopen(fd,'wb') as stream:
                    stream.write(src.read_bytes())
                os.chmod(tmp_name, 0o644)
                os.replace(tmp_name,dest)
            finally:
                if os.path.exists(tmp_name):
                    os.unlink(tmp_name)
    return target,manifest

def main(argv=None):
    parser=argparse.ArgumentParser(description='Validate canonical event XLSX and generate independent app/materials.')
    parser.add_argument('workbook',type=Path)
    parser.add_argument('--output',type=Path,help='Default: generated/<event.slug>; never use one directory for multiple events.')
    parser.add_argument('--template',type=Path,help='Optional runtime template directory.')
    parser.add_argument('--validate-only',action='store_true')
    args=parser.parse_args(argv)
    try:
        if args.validate_only:
            config,assets=load_config(args.workbook)
            print(f"配置有效：{config['event']['slug']}，{len(config['agenda'])} 条日程，{len(assets)} 个公共素材")
            if config['branding'].get('heroWarning'):print('提示：'+config['branding']['heroWarning'])
        else:
            path,manifest=generate(args.workbook,args.output,args.template)
            print(f'已生成：{path}')
            print(f"{len(manifest['artifacts'])} 个文件；配置 SHA-256 {manifest['configSha256']}")
            for warning in manifest['warnings']:
                print('提示：'+warning)
    except ConfigError as exc:
        print('配置错误：'+str(exc),file=sys.stderr)
        return 2
    return 0

if __name__=='__main__':
    raise SystemExit(main())
